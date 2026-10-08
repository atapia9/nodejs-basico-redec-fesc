#!/usr/bin/env python3
"""Verifica la fidelidad del Markdown del repositorio frente a los .docx fuente.

Comprueba: texto (0 párrafos faltantes), conteos estructurales, fechas, enlaces
del Anexo 12 y que el código de ejemplo no se modificó.

Requiere: Python 3.8+ y python-docx (pip install python-docx).
Requiere también los .docx en fuentes/ (esa carpeta no se publica en GitHub).

Uso:
    python3 scripts/verificar_fidelidad.py [--repo RAIZ] [--fuentes CARPETA]

Código de salida: 0 si no hay fallas, 1 si alguna comprobación falla.
"""
import argparse
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit('Falta python-docx. Instálalo con: pip install python-docx')

MANUAL = 'Manual_NodeJS_Basico_ilustrado_v3.docx'
FICHAS = 'Fichas_Videos_NodeJS.docx'
EVALUACION = 'Evaluacion_Diagnostica_Final_NodeJS.docx'

PLAYLIST = 'https://youtube.com/playlist?list=PLUMolTkur2Cc&si=RUIZHDWXTIkthQn9'
FIGURAS = ['1.1', '1.2', '2.1', '2.2', '2.3', '3.1', '3.2', '4.1', '4.2', '5.1', '5.2']
CALENDARIO = [
    (1, 'Lunes 19 de octubre de 2026'),
    (2, 'Martes 20 de octubre de 2026'),
    (3, 'Miércoles 21 de octubre de 2026'),
    (4, 'Jueves 22 de octubre de 2026'),
    (5, 'Viernes 23 de octubre de 2026'),
]
# Líneas que solo existen en la portada del archivo independiente de fichas.
DIFERENCIAS_FICHAS = {
    'Videos de apoyo — Node.js Básico',
    'Fichas por sesión · REDEC-UNAM / FESC · 19 al 23 de octubre de 2026',
}
FILL_EX = 'F2F2F2'
INSTRUCTOR = 'SOLO PARA EL INSTRUCTOR'
# Mes de las fechas anteriores del curso; se escribe partido para que este archivo no se detecte a sí mismo.
MES_OBSOLETO = 'septi' + 'embre'


# ------------------------------------------------------------------ normalización
def norm(s):
    s = s.replace('\u00a0', ' ')
    s = re.sub(r'^\s*•\s*', '', s)
    return re.sub(r'\s+', ' ', s).strip()


def clean_code(text):
    lines = ['' if not ln.strip() else ln for ln in text.split('\n')]
    while lines and lines[-1] == '':
        lines.pop()
    while lines and lines[0] == '':
        lines.pop(0)
    return '\n'.join(lines)


def inline_plain(s):
    s = re.sub(r'!\[[^\]]*\]\([^)]*\)', '', s)
    s = re.sub(r'\[((?:\\.|[^\]\\])*)\]\([^)]*\)', r'\1', s)
    s = re.sub(r'(?<!\\)\*+', '', s)
    s = re.sub(r'\\(.)', r'\1', s)
    return norm(s)


def split_row(line):
    cells = re.split(r'(?<!\\)\|', line.strip())
    if cells and cells[0].strip() == '':
        cells = cells[1:]
    if cells and cells[-1].strip() == '':
        cells = cells[:-1]
    return cells


def md_units(text):
    """Unidades de texto normalizadas de un .md (párrafos, celdas, líneas de código)."""
    units = []
    in_fence = False
    for raw in text.split('\n'):
        line = raw.rstrip()
        if line.startswith('```'):
            in_fence = not in_fence
            continue
        if in_fence:
            if line.strip():
                units.append(norm(line))
            continue
        line = re.sub(r'^\s*(>\s?)+', '', line)
        es_encabezado = bool(re.match(r'^#{1,6}\s+', line))
        line = re.sub(r'^#{1,6}\s+', '', line)
        if not line.strip() or re.match(r'^\s*!\[', line):
            continue
        if line.strip().startswith('|'):
            if set(line.replace('|', '').replace(':', '').replace(' ', '')) <= {'-'}:
                continue
            for cell in split_row(line):
                for part in re.split(r'<br\s*/?>', cell):
                    u = inline_plain(part)
                    if u:
                        units.append(u)
            continue
        if not es_encabezado:
            line = re.sub(r'^\s*([-+]|\d+[.)])\s+', '', line)
        u = inline_plain(line)
        if u:
            units.append(u)
    return units


def docx_units(path):
    """Unidades de texto normalizadas de un .docx (párrafos y párrafos de celdas)."""
    doc = Document(str(path))
    units = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn('w:p'):
            for part in Paragraph(el, doc).text.split('\n'):
                if part.strip():
                    units.append(norm(part))
        elif el.tag == qn('w:tbl'):
            for row in Table(el, doc).rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        for part in p.text.split('\n'):
                            if part.strip():
                                units.append(norm(part))
    return units


def fill_of(tc):
    p = tc.tcPr
    s = p.find(qn('w:shd')) if p is not None else None
    return s.get(qn('w:fill')) if s is not None else None


def docx_blocks(path):
    """Devuelve (bloques_de_codigo, titulos_de_ejercicio) de un .docx."""
    doc = Document(str(path))
    codes, exercises = [], []
    for t in doc.tables:
        if len(t.columns) != 1:
            continue
        for r in t.rows:
            c = r.cells[0]
            f = fill_of(c._tc)
            if f == FILL_EX or c.text.strip().lower().startswith('ejercicio'):
                exercises.append(c.paragraphs[0].text.strip())
            else:
                codes.append(clean_code('\n'.join(p.text for p in c.paragraphs)))
    return codes, exercises


def docx_figure_captions(path):
    doc = Document(str(path))
    els = list(doc.element.body.iterchildren())
    caps = []
    for i, el in enumerate(els):
        if el.tag == qn('w:p') and el.findall('.//' + qn('w:drawing')):
            caps.append(Paragraph(els[i + 1], doc).text.strip())
    return caps


def docx_hyperlinks(path):
    z = zipfile.ZipFile(str(path))
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
    rel_urls = {}
    for m in re.finditer(r'<Relationship\b([^>]*)/?>', rels):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        if attrs.get('TargetMode') == 'External' and attrs.get('Target', '').startswith('http'):
            rel_urls[attrs['Id']] = attrs['Target'].replace('&amp;', '&')
    xml = z.read('word/document.xml').decode('utf-8')
    used = re.findall(r'<w:hyperlink\b[^>]*r:id="(rId\d+)"', xml)
    return set(rel_urls.values()), [rel_urls[r] for r in used if r in rel_urls]


def fenced_blocks(text):
    out, cur, lang, inside = [], [], None, False
    for line in text.split('\n'):
        if line.startswith('```'):
            if inside:
                out.append((lang, '\n'.join(cur)))
                cur, inside = [], False
            else:
                lang, inside = line[3:].strip(), True
            continue
        if inside:
            cur.append(line)
    return out


# ------------------------------------------------------------------ comprobaciones
class Reporte:
    def __init__(self):
        self.filas = []

    def add(self, grupo, nombre, estado, detalle=''):
        self.filas.append((grupo, nombre, estado, detalle))

    def imprimir(self):
        w1 = max(len(f[1]) for f in self.filas)
        print('\n%-3s %-*s  %-6s  %s' % ('#', w1, 'Comprobación', 'Estado', 'Detalle'))
        print('-' * (w1 + 40))
        for n, (g, nombre, estado, detalle) in enumerate(self.filas, 1):
            print('%-3d %-*s  %-6s  %s' % (n, w1, nombre, estado, detalle))
        fallas = [f for f in self.filas if f[2] == 'FALLA']
        avisos = [f for f in self.filas if f[2] == 'AVISO']
        print('-' * (w1 + 40))
        print('Resultado: %d comprobaciones, %d fallas, %d avisos.' % (len(self.filas), len(fallas), len(avisos)))
        return 1 if fallas else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--repo', default=str(Path(__file__).resolve().parent.parent))
    ap.add_argument('--fuentes', default=None)
    args = ap.parse_args()
    repo = Path(args.repo)
    fuentes = Path(args.fuentes) if args.fuentes else repo / 'fuentes'
    for n in (MANUAL, FICHAS, EVALUACION):
        if not (fuentes / n).exists():
            sys.exit('No se encuentra %s en %s' % (n, fuentes))
    rep = Reporte()

    docs = sorted((repo / 'docs').glob('*.md'))
    docs_sin_notas = [p for p in docs if p.name != 'NOTAS_DE_REVISION.md']
    readme = repo / 'README.md'
    ev_pub = repo / 'evaluacion' / 'evaluacion-diagnostica-final.md'
    ev_clave = repo / 'evaluacion' / '_clave-instructor.md'
    videos = repo / 'docs' / '12-videos-de-apoyo.md'

    # ---------- 1. Texto
    corpus = Counter()
    for p in [readme] + docs_sin_notas:
        corpus.update(md_units(p.read_text(encoding='utf-8')))
    esperado = Counter(docx_units(fuentes / MANUAL))
    faltan = []
    for u, n in esperado.items():
        if corpus[u] < n:
            faltan.append((u, n - corpus[u]))
    total = sum(esperado.values())
    rep.add('texto', 'Texto del manual: párrafos del .docx ausentes en el .md',
            'OK' if not faltan else 'FALLA',
            '%d unidades verificadas, %d faltantes' % (total, sum(n for _, n in faltan)))
    for u, n in faltan[:15]:
        print('  faltante (%dx): %s' % (n, u[:140]))

    # Fichas independientes frente al Anexo 12
    corp_v = Counter(md_units(videos.read_text(encoding='utf-8')))
    esp_f = Counter(docx_units(fuentes / FICHAS))
    falt_f = [u for u, n in esp_f.items() if corp_v[u] < n and u not in DIFERENCIAS_FICHAS]
    rep.add('texto', 'Fichas_Videos_NodeJS.docx frente al Anexo 12 (docs/12)',
            'OK' if not falt_f else 'FALLA',
            '%d unidades, %d faltantes, %d líneas de portada propias de las fichas'
            % (sum(esp_f.values()), len(falt_f), len(DIFERENCIAS_FICHAS)))
    for u in falt_f[:10]:
        print('  faltante en fichas: %s' % u[:140])

    # Evaluación
    ev_text = ev_pub.read_text(encoding='utf-8')
    ev_corpus = Counter(md_units(ev_text))
    clave_local = ev_clave.exists()
    if clave_local:
        ev_corpus.update(md_units(ev_clave.read_text(encoding='utf-8')))
    ev_units = docx_units(fuentes / EVALUACION)
    if INSTRUCTOR in ev_units:
        k = ev_units.index(INSTRUCTOR)
        parte_alumno, parte_clave = ev_units[:k], ev_units[k:]
    else:
        parte_alumno, parte_clave = ev_units, []
    exigidas = Counter(parte_alumno + (parte_clave if clave_local else []))
    falt_e = [u for u, n in exigidas.items() if ev_corpus[u] < n]
    detalle = '%d unidades, %d faltantes' % (sum(exigidas.values()), len(falt_e))
    if not clave_local:
        detalle += '; clave del instructor no disponible localmente (%d unidades omitidas)' % len(parte_clave)
    rep.add('texto', 'Evaluación diagnóstica: hoja del alumno y clave',
            'OK' if not falt_e else 'FALLA', detalle)
    for u in falt_e[:10]:
        print('  faltante en evaluación: %s' % u[:140])
    fuga = [u for u in parte_clave if u in Counter(md_units(ev_text)) and u not in Counter(parte_alumno)]
    rep.add('texto', 'La hoja pública del alumno no contiene la clave de respuestas',
            'OK' if not fuga else 'FALLA', '%d líneas de la clave halladas en la hoja pública' % len(fuga))

    # ---------- 2. Conteos
    sesiones = sorted(p for p in (repo / 'docs').glob('sesion-*.md'))
    ok_s = len(sesiones) == 5 and all(
        re.match(r'^# \d+\. Sesión \d — ', p.read_text(encoding='utf-8').split('\n')[0]) for p in sesiones)
    rep.add('conteos', 'Sesiones (docs/sesion-N-*.md)', 'OK' if ok_s else 'FALLA', '%d archivos' % len(sesiones))

    md_figs = []
    for p in docs_sin_notas:
        t = p.read_text(encoding='utf-8')
        for m in re.finditer(r'!\[[^\]]*\]\(\.\./assets/figuras/fig-(\d)-(\d)\.png\)\n\n\*(Figura \d\.\d: [^\n]*?)\*\n', t):
            md_figs.append(('%s.%s' % (m.group(1), m.group(2)), m.group(3).replace('\\', '')))
    docx_caps = [norm(c) for c in docx_figure_captions(fuentes / MANUAL)]
    ok_f = (sorted(i for i, _ in md_figs) == sorted(FIGURAS)
            and [norm(c) for _, c in md_figs] == docx_caps
            and all(c.startswith('Figura %s:' % i) for i, c in md_figs))
    rep.add('conteos', 'Figuras con imagen y pie idéntico al manual', 'OK' if ok_f else 'FALLA',
            '%d de %d figuras con pie' % (len(md_figs), len(FIGURAS)))
    png = [f for f in FIGURAS if (repo / 'assets' / 'figuras' / ('fig-%s.png' % f.replace('.', '-'))).exists()]
    svg = [f for f in FIGURAS if (repo / 'assets' / 'figuras' / ('fig-%s.svg' % f.replace('.', '-'))).exists()]
    rep.add('conteos', 'Archivos de figura PNG', 'OK' if len(png) == 11 else 'FALLA', '%d de 11' % len(png))
    rep.add('conteos', 'Archivos de figura SVG (imagen principal)',
            'OK' if len(svg) == 11 else 'AVISO',
            '%d de 11%s' % (len(svg), '' if len(svg) == 11 else ' (pendiente: no se contó con Figuras_NodeJS_SVG.zip)'))

    codes, ejercicios = docx_blocks(fuentes / MANUAL)
    md_codes = []
    for p in docs_sin_notas:
        md_codes += [b for _, b in fenced_blocks(p.read_text(encoding='utf-8'))]
    falt_c = list((Counter(codes) - Counter(md_codes)).elements())
    rep.add('conteos', 'Bloques de código del .docx presentes (texto exacto, con sangrías)',
            'OK' if not falt_c and len(md_codes) == len(codes) else 'FALLA',
            '%d en el .docx, %d en el .md, %d sin coincidencia exacta' % (len(codes), len(md_codes), len(falt_c)))
    for b in falt_c[:5]:
        print('  bloque sin coincidencia exacta: %r' % b[:80])

    md_ej = 0
    for p in docs_sin_notas:
        md_ej += len(re.findall(r'^> \*\*Ejercicio', p.read_text(encoding='utf-8'), re.M))
    rep.add('conteos', 'Cuadros «Ejercicio…» presentes', 'OK' if md_ej == len(ejercicios) else 'FALLA',
            '%d en el .docx, %d en el .md' % (len(ejercicios), md_ej))

    reactivos = sorted(set(int(n) for n in re.findall(r'^\*\*(\d+)\\?\. ', ev_text, re.M)))
    rep.add('conteos', 'Reactivos de la evaluación', 'OK' if reactivos == list(range(1, 11)) else 'FALLA',
            '%d reactivos' % len(reactivos))
    if clave_local:
        filas = re.findall(r'^\| (\d+) \| ([A-D]) \|', ev_clave.read_text(encoding='utf-8'), re.M)
        rep.add('conteos', 'Clave del instructor (archivo local)', 'OK' if len(filas) == 10 else 'FALLA',
                '%d respuestas' % len(filas))
    ev_codes, _ = docx_blocks(fuentes / EVALUACION)
    ev_md_codes = [b for _, b in fenced_blocks(ev_text)]
    rep.add('conteos', 'Bloques de código de la evaluación',
            'OK' if not (Counter(ev_codes) - Counter(ev_md_codes)) else 'FALLA',
            '%d en el .docx, %d en el .md' % (len(ev_codes), len(ev_md_codes)))

    v = videos.read_text(encoding='utf-8')
    subtemas = re.findall(r'^### (\d)\.(\d) · ', v, re.M)
    por_sesion = Counter(s for s, _ in subtemas)
    rep.add('conteos', 'Subtemas del Anexo 12 (5 sesiones)',
            'OK' if len(subtemas) == 19 and len(re.findall(r'^## Sesión \d: ', v, re.M)) == 5 else 'FALLA',
            '%d subtemas (%s)' % (len(subtemas), ', '.join('S%s=%d' % (k, por_sesion[k]) for k in sorted(por_sesion))))
    n_tablas = len(re.findall(r'^\| \*\*Tipo\*\* \| \*\*Video \(clic para abrir\)\*\* \|$', v, re.M))
    n_ojo = len(re.findall(r'^\*Ojo: ', v, re.M))
    n_cierre = len(re.findall(r'^\*\*Cierre:\*\*', v, re.M))
    n_act = len(re.findall(r'^Actividad \d', v, re.M))
    rep.add('conteos', 'Fichas: tablas Tipo/Video, notas «Ojo», cierres y actividades',
            'OK' if (n_tablas == 19 and n_cierre == 5 and n_act == 5) else 'FALLA',
            '%d tablas, %d notas Ojo, %d cierres, %d actividades' % (n_tablas, n_ojo, n_cierre, n_act))

    # ---------- 3. Fechas
    pat = re.compile(MES_OBSOLETO, re.I)
    ext = {'.md', '.js', '.mjs', '.json', '.jsonc', '.yml', '.yaml', '.txt', '.ejs', '.py', '.toml', '.cjs'}
    ignorar = {'.git', 'fuentes', 'node_modules'}
    hallazgos = []
    for p in repo.rglob('*'):
        if p.is_file() and not (set(p.relative_to(repo).parts) & ignorar) and (p.suffix in ext or p.name.startswith('.')):
            try:
                if pat.search(p.read_text(encoding='utf-8')):
                    hallazgos.append(str(p.relative_to(repo)))
            except (UnicodeDecodeError, OSError):
                pass
    rep.add('fechas', 'Ninguna mención a «%s» en el repositorio' % MES_OBSOLETO, 'OK' if not hallazgos else 'FALLA',
            ', '.join(hallazgos) if hallazgos else 'sin menciones')
    ok_cal = True
    for n, fecha in CALENDARIO:
        f = next((p for p in sesiones if p.name.startswith('sesion-%d-' % n)), None)
        if f is None or fecha not in f.read_text(encoding='utf-8'):
            ok_cal = False
    ok_cal = ok_cal and 'Del 19 al 23 de octubre de 2026' in readme.read_text(encoding='utf-8')
    rep.add('fechas', 'Calendario: lunes 19 a viernes 23 de octubre de 2026', 'OK' if ok_cal else 'FALLA',
            '; '.join(f for _, f in CALENDARIO))

    # ---------- 4. Enlaces
    rels, usados = docx_hyperlinks(fuentes / MANUAL)
    rels_f, _ = docx_hyperlinks(fuentes / FICHAS)
    md_urls = re.findall(r'\]\((https?://[^)]+)\)', v)
    ok_l = set(md_urls) == rels == set(usados) == rels_f
    rep.add('enlaces', 'URLs del Anexo 12: conjunto del .md idéntico al del .docx',
            'OK' if ok_l else 'FALLA',
            '%d únicas en el .md, %d en word/_rels, %d en las fichas; %d hipervínculos en el .docx y %d en el .md'
            % (len(set(md_urls)), len(rels), len(rels_f), len(usados), len(md_urls)))
    if not ok_l:
        print('  solo en .md:', sorted(set(md_urls) - rels)[:5], '| solo en .docx:', sorted(rels - set(md_urls))[:5])
    rep.add('enlaces', 'Enlace a la lista de reproducción (no listada) al inicio del Anexo 12',
            'OK' if ('](%s)' % PLAYLIST) in v.split('## Sesión 1')[0] else 'FALLA', PLAYLIST)

    # ---------- 5. Código de ejemplo
    ejemplos = [p for p in (repo / 'ejemplos').rglob('*') if p.is_file()]
    no_fieles = []
    for p in ejemplos:
        cont = p.read_text(encoding='utf-8').rstrip('\n')
        if not any(cont in b for b in codes):
            no_fieles.append(str(p.relative_to(repo)))
    rep.add('código', 'Archivos de ejemplos/ sin modificar respecto de los bloques del manual',
            'OK' if not no_fieles else 'FALLA', '%d archivos, %d sin coincidencia' % (len(ejemplos), len(no_fieles)))
    for n in no_fieles:
        print('  sin coincidencia:', n)

    sys.exit(rep.imprimir())


if __name__ == '__main__':
    main()
